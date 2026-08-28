"""Build the audited research article DOCX for the OneB Market project."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/article/oneb_market_audited_research_article.docx")
NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(89, 89, 89)


def set_font(run, name="Aptos", size=10.5, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
        col.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9, bold=True, color=NAVY)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_font(run, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(text)
        set_font(r)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size={1: 15, 2: 12.5}.get(level, 11), bold=True, color=NAVY if level == 1 else BLUE)
    return p


def add_page_number(footer):
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("OneB Market | Research audit | ")
    set_font(r, size=8.5, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_page_number(section.footer)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    for name, size, color in (("Heading 1", 15, NAVY), ("Heading 2", 12.5, BLUE)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    # Editorial-cover opening for a research report.
    doc.add_paragraph().paragraph_format.space_after = Pt(52)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ARTIGO DE PESQUISA - VERSÃO AUDITADA")
    set_font(r, size=10, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(15)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sinais Técnicos Condicionados a Regime em Ações dos EUA")
    set_font(r, name="Aptos Display", size=25, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(7)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Uma auditoria reproduzível de dados, validação temporal e prontidão operacional no OneB Market")
    set_font(r, size=13, color=GRAY)
    p.paragraph_format.space_after = Pt(38)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Base auditada em 21 de agosto de 2026")
    set_font(r, size=10.5, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Projeto OneB Market | Rascunho para submissão e revisão técnica")
    set_font(r, size=9.5, italic=True, color=GRAY)
    doc.add_page_break()

    heading(doc, "Resumo")
    add_para(doc, "Este artigo apresenta uma auditoria do pipeline de pesquisa quantitativa do OneB Market e investiga se indicadores técnicos carregam informação preditiva em ações norte-americanas quando condicionados ao regime de mercado. A base auditada contém 10.368 observações símbolo-data de 24 ativos, com 40 features utilizáveis, retornos futuros em múltiplos horizontes e separação cronológica entre treino, validação e teste. A auditoria encontrou zero duplicatas símbolo-data e zero campos ausentes na tabela de pesquisa. O estudo adota validação walk-forward, embargo temporal, comparação com baseline, custos de 14 bps por giro e verificações explícitas de estabilidade temporal e circularidade. A maior parte das formulações testadas não sobreviveu ao holdout. O resultado mais consistente é um candidato de duas features - volatilidade anualizada e ATR percentual - aplicado apenas em dias BEAR, com AUC de aproximadamente 0,53 em janelas de cinco e dez anos. Esse efeito é pequeno e não é tratado como uma estratégia automatizável. A principal contribuição é metodológica: demonstrar um fluxo reprodutível que prioriza a rejeição de falsos positivos e separa evidência estatística de prontidão operacional.")
    add_para(doc, "Palavras-chave: finanças quantitativas; validação walk-forward; regime de mercado; auditoria de dados; aprendizado de máquina; reprodutibilidade.")

    heading(doc, "I. Introdução")
    add_para(doc, "Sistemas de mercado frequentemente confundem visualização, backtest e previsão. Um retorno histórico elevado, por si só, não demonstra poder preditivo: pode refletir ajuste excessivo, dependência temporal, custos subestimados ou uma seleção de hipóteses feita depois de observar os dados. O OneB Market foi construído como plataforma de monitoramento, alertas, pesquisa e simulação, sem execução de ordens. Essa separação permite que a pesquisa seja auditada antes de ser acoplada a qualquer fluxo de decisão.")
    add_para(doc, "A questão investigada não é se indicadores técnicos produzem uma recomendação de compra ou venda. A pergunta é mais restrita: existe informação cross-sectional estável em features técnicas, especialmente após separar regimes de mercado e reservar dados futuros para validação? A resposta atual é mista: há um efeito pequeno e repetido em BEAR, mas não há evidência suficiente para autonomia operacional.")

    heading(doc, "II. Auditoria dos dados e reprodutibilidade")
    add_para(doc, "A auditoria combinou a inspeção do dataset de pesquisa, o gate de confiabilidade dos candles, os artefatos de validação e o histórico de versões do repositório. Os resultados abaixo descrevem o estado observado em 21 de agosto de 2026; datas diferentes entre artefatos são reportadas como risco de governança, não suavizadas.")
    add_table(doc, ["Item", "Evidência auditada", "Avaliação"], [
        ["Painel de pesquisa", "10.368 linhas; 24 símbolos; 40 features utilizáveis; 432 datas cross-sectional", "Aprovado para pesquisa"],
        ["Janela efetiva", "22/10/2024 a 15/07/2026; labels futuros exigem truncamento no fim", "Coerente com o horizonte"],
        ["Integridade tabular", "0 duplicatas símbolo-data; 0 células vazias; splits: 7.248 treino, 1.560 validação, 1.560 teste", "Aprovado"],
        ["Candles EOD", "24/24 ativos aprovados no gate de 13/08; 501 candles por ativo; fonte Tiingo", "Aprovado, mas gate precisa ser renovado"],
        ["Cache recente", "Candles Tiingo de 2 anos atualizados em 21/08", "Atualizado"],
        ["Macro", "DXY, Treasury 10Y e VIX aprovados; última observação no gate: 07/08", "Revalidar antes de nova rodada"],
    ], [1800, 4700, 2860])
    add_para(doc, "O principal ponto de atenção não é um defeito encontrado na tabela de pesquisa, mas a defasagem entre os artefatos: o gate formal é de 13 de agosto, enquanto caches e simulações foram atualizados em 21 de agosto. Para um artigo, cada resultado deve carregar um manifesto com versão do dataset, intervalo de datas, provedor, hash do código e horário de geração. Essa é a melhoria de maior retorno científico imediato.")

    heading(doc, "III. Metodologia")
    add_bullets(doc, [
        "Dados: candles diários ajustados, benchmark QQQ e universo de 24 ações líquidas dos EUA; as features usam somente informação disponível na data da observação.",
        "Alvo: retorno futuro de cinco dias e ranking cross-sectional acima/abaixo da mediana diária, apropriado para avaliar seleção relativa de ativos.",
        "Validação: folds temporais walk-forward com embargo/purging; a seleção de features ocorre no treino, nunca no holdout.",
        "Controles: baseline de classe majoritária, AUC, Brier score, custos de ida e volta de 14 bps, avaliação por regime e estabilidade entre períodos.",
        "Antiviés: correção de múltiplas comparações, checagem de circularidade entre features e labels e retenção explícita de experimentos negativos.",
    ])
    add_para(doc, "A nova rotina regime_signal_evidence.py estende essa disciplina com bootstrap por blocos de data e teste de permutação que embaralha labels apenas dentro de cada pregão. A rotina está testada unitariamente, porém a execução final com dados de mercado deve ser refeita em ambiente com conectividade e armazenada como artefato do experimento. Portanto, seus resultados numéricos ainda não são usados neste artigo.")

    heading(doc, "IV. Resultados")
    add_table(doc, ["Hipótese / artefato", "Resultado", "Leitura correta"], [
        ["Modelo probabilístico geral", "Acurácia holdout 48,99%; baseline 51,49%", "Não promover"],
        ["Features técnicas pooled", "Correlações fracas; diversos sinais perdem estabilidade temporal", "Não há edge geral demonstrado"],
        ["Sinal BEAR: volatilidade + ATR%", "AUC 0,5321 (5 anos) e 0,5329 (10 anos)", "Efeito pequeno, mas replicado"],
        ["Estratégia cross-sectional", "+23,836% em 13 meses; drawdown -24,447%; t mensal 0,872", "Exploratória; evidência insuficiente"],
        ["Technical Edge, replay recente", "+30,724% em 30 períodos; drawdown -13,893%; 14 bps por giro", "Amostra curta e seleção adaptativa; pesquisa apenas"],
        ["Prontidão para automação", "Long e short falharam todos os quatro folds", "Aprovação humana obrigatória"],
    ], [2350, 2500, 4510])
    add_para(doc, "O achado BEAR é interessante porque a repetição entre duas janelas históricas distintas é mais informativa que uma correlação agregada. Ainda assim, AUC próximo de 0,53 não sustenta alocação autônoma. O resultado de estratégia recente também não deve ser convertido em claim de desempenho: ele contém apenas 30 períodos, turnover médio de 1,62 e uma regra de features selecionadas dentro do procedimento. A seção de resultados, portanto, enfatiza a incerteza e não o retorno acumulado.")

    heading(doc, "V. Estado atual e atualizações recentes")
    add_bullets(doc, [
        "21/08: validação rápida do simulador confirmou dados OHLCV locais para AAPL, MSFT, NVDA e SNAP, mas não autorizou operação: can_trade_now=false.",
        "21/08: validação cross-sectional foi atualizada com 10.368 linhas e 40 features utilizáveis; o relatório marcou status PASS como integridade do experimento, não como aprovação de trading.",
        "21/08: uma simulação do Technical Edge foi registrada. Ela é útil como hipótese operacional, mas precisa de uma divisão externa e fixa antes de qualquer conclusão.",
        "20/08: o paper simulator registrou eventos recentes; ainda é uma carteira fictícia e não substitui um track record forward ao vivo.",
        "10/08: o relatório de automação concluiu HUMAN_APPROVAL_REQUIRED para os lados long e short.",
    ])

    heading(doc, "VI. Próximas melhorias priorizadas")
    add_table(doc, ["Prioridade", "Melhoria", "Critério de conclusão"], [
        ["P0 - governança", "Gerar um manifesto por experimento: hash Git, versão/datas do dataset, provedor, parâmetros, seed e checksums.", "Qualquer tabela do artigo pode ser reproduzida a partir de um único comando."],
        ["P0 - atualização", "Executar data_reliability_gate antes de cada rodada e bloquear resultados com macro/candles defasados.", "Gate recente e anexo ao artefato de pesquisa."],
        ["P0 - inferência", "Rodar o relatório BEAR com bootstrap por data e teste de permutação; publicar IC 95%, Brier e p-valor.", "Efeito medido com incerteza, não só AUC pontual."],
        ["P1 - generalização", "Replicar a hipótese pré-especificada num segundo universo independente e em mais episódios BEAR.", "Sinal preservado fora do universo original ou resultado negativo documentado."],
        ["P1 - economia", "Avaliar custos, turnover, capacidade, drawdown e benchmark passivo no mesmo holdout.", "Decisão de pesquisa separada da decisão operacional."],
        ["P2 - dados", "Acumular notícias timestamped e considerar opções/order flow somente com fonte licenciada e cobertura histórica adequada.", "Novo dataset point-in-time com auditoria equivalente."],
    ], [1150, 4050, 4160])

    heading(doc, "VII. Conclusão")
    add_para(doc, "O OneB Market já é mais maduro como plataforma de pesquisa do que como sistema de previsão. A base de dados pesquisada é limpa e o processo contém proteções relevantes contra vazamento, sobreajuste e narrativas de retorno. O saldo científico atual é deliberadamente conservador: a maioria dos sinais testados não generaliza, enquanto o candidato condicionado a BEAR apresenta uma repetição pequena, ainda insuficiente para automação. Essa conclusão é publicável porque transforma resultado negativo, controle metodológico e governança de dados em contribuição técnica verificável.")
    add_para(doc, "A recomendação é congelar agora uma hipótese primária, renovar o gate de dados, produzir os intervalos de confiança por bloco temporal e executar uma replicação externa. Só depois deve ser avaliada uma camada de uso operacional, sempre com aprovação humana enquanto os critérios de prontidão continuarem reprovados.")

    heading(doc, "Referências e artefatos internos", 1)
    refs = [
        "[1] OneB Market, ResearchDataset v1: especificação e resumo de geração, docs/ml/research-dataset-v1.md, 2026.",
        "[2] OneB Market, Data Phase Findings: auditoria de features e validação por regime, docs/data_phase_findings.md, 2026.",
        "[3] OneB Market, data/data_reliability_gate.json, gerado em 13 ago. 2026.",
        "[4] OneB Market, data/cross_sectional_strategy_validation.json e data/technical_edge_simulation.json, atualizados em 21 ago. 2026.",
        "[5] OneB Market, data/automation_readiness_report.json e data/probability_model_history.json, 2026.",
        "[6] M. López de Prado, Advances in Financial Machine Learning. Hoboken, NJ, USA: Wiley, 2018. (referência metodológica para purging e embargo).",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(ref)
        set_font(r, size=9)

    doc.core_properties.title = "Sinais Técnicos Condicionados a Regime em Ações dos EUA"
    doc.core_properties.subject = "Artigo de pesquisa auditado do OneB Market"
    doc.core_properties.author = "OneB Market"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
